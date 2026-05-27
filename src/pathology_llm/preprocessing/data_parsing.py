from pathlib import Path


def load_reports_from_excel(
    path: str | Path,
    id_column: str = "Number",
    text_column: str = "Final Diagnosis",
) -> list[str]:
    """Load report texts from an Excel file.

    Expected columns:
    - Number: integer report identifier
    - Final Diagnosis: free-text report content
    """
    try:
        from openpyxl import load_workbook  # type: ignore[import-not-found]
    except ImportError as exc:
        raise RuntimeError(
            "Excel parsing requires openpyxl. Install it with: pip install openpyxl"
        ) from exc

    workbook = load_workbook(filename=Path(path), read_only=True, data_only=True)
    sheet = workbook.active

    header_cells = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_cells:
        raise ValueError(f"Excel file has no header row: {path}")

    headers = [
        str(value).strip() if value is not None else "" for value in header_cells
    ]
    if id_column not in headers or text_column not in headers:
        raise ValueError(
            f"Excel file must contain columns '{id_column}' and '{text_column}'. Found: {headers}"
        )

    id_index = headers.index(id_column)
    text_index = headers.index(text_column)

    reports: list[str] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        report_id = row[id_index] if id_index < len(row) else None
        report_text = row[text_index] if text_index < len(row) else None

        if report_id is None and (
            report_text is None or str(report_text).strip() == ""
        ):
            continue

        if report_id is not None and not isinstance(report_id, int):
            if isinstance(report_id, float) and report_id.is_integer():
                report_id = int(report_id)
            else:
                raise ValueError(
                    f"Invalid report ID in column '{id_column}': {report_id!r}"
                )

        if report_text is None:
            continue

        text = str(report_text).strip()
        if text:
            reports.append(text)

    if not reports:
        raise ValueError(f"No report texts found in column '{text_column}' from {path}")

    return reports
