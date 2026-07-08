from collections.abc import Iterable
from pathlib import Path
from typing import Any
from openpyxl.worksheet.table import Table
import yaml
import argparse
import logging

from postprocessing.predictions import (
    load_prediction_records,
    normalize_case_id,
    to_excel_value,
)

logger = logging.getLogger(__name__)


def iter_header_values(values: Iterable[Any]) -> list[str]:
    return [str(value).strip() if value is not None else "" for value in values]


def find_header_index(headers: list[str], target: str) -> int | None:
    if target in headers:
        return headers.index(target)

    target_folded = target.strip().casefold()
    for index, header in enumerate(headers):
        if header.strip().casefold() == target_folded:
            return index
    return None


def normalize_value_for_match(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value).casefold()
    text = str(value).strip()
    if not text:
        return None
    return text.casefold()


def normalize_containers(value: Any) -> list[dict[str, Any]]:
    if value is None:
        return []

    if isinstance(value, list):
        items = value
    else:
        items = [value]

    normalized: list[dict[str, Any]] = []
    for item in items:
        if isinstance(item, dict):
            normalized.append(item)
            continue
        if item is None:
            continue
        normalized.append({"label": str(item)})

    return normalized


def resolve_container_label(container: dict[str, Any] | None) -> Any:
    if not container:
        return None
    label = container.get("label")
    return to_excel_value(label)


def resolve_container_field(container: dict[str, Any] | None, field: str) -> Any:
    if not container:
        return None
    return to_excel_value(container.get(field))


def extract_primary_diagnosis_field(record: dict[str, Any], field: str) -> Any:
    """Extract primary diagnosis field from nested or flat structure.

    Handles new nested structure: valid_primary_diagnoses -> top_1 -> field
    Falls back to old flat structure: field (for backward compatibility)
    """
    if "valid_primary_diagnoses" in record and isinstance(
        record["valid_primary_diagnoses"], dict
    ):
        top_1 = record["valid_primary_diagnoses"].get("top_1")
        if top_1 and isinstance(top_1, dict):
            return to_excel_value(top_1.get(field))
    # Fallback to flat structure
    return to_excel_value(record.get(field))


def extract_container_diagnosis_field(
    container: dict[str, Any] | None, field: str
) -> Any:
    """Extract container diagnosis field from nested or flat structure.

    Handles new nested structure: valid_diagnoses -> top_1 -> field
    Falls back to old flat structure: field (for backward compatibility)
    """
    if not container:
        return None
    if "valid_diagnoses" in container and isinstance(
        container["valid_diagnoses"], dict
    ):
        top_1 = container["valid_diagnoses"].get("top_1")
        if top_1 and isinstance(top_1, dict):
            return to_excel_value(top_1.get(field))
    # Fallback to flat structure
    return to_excel_value(container.get(field))


def select_container_for_row(
    record: dict[str, Any],
    excel_container_value: Any,
) -> dict[str, Any] | None:
    containers = normalize_containers(record.get("containers"))
    if not containers and "container" in record:
        containers = normalize_containers(record.get("container"))
    if not containers:
        return None

    container_by_label = {}
    for container in containers:
        label = container.get("label")
        if label is None:
            continue
        label_normalized = str(label).strip().casefold()
        if label_normalized:
            container_by_label[label_normalized] = container

    excel_label = (
        str(excel_container_value).strip().casefold() if excel_container_value else None
    )
    if excel_label and excel_label in container_by_label:
        return container_by_label[excel_label]

    return None


def resolve_fallback_container_label(record: dict[str, Any]) -> Any:
    containers = normalize_containers(record.get("containers"))
    if not containers and "container" in record:
        containers = normalize_containers(record.get("container"))

    for container in containers:
        label = resolve_container_label(container)
        if label not in {None, ""}:
            return label
    return None


def resolve_fallback_container_field(record: dict[str, Any], field: str) -> Any:
    containers = normalize_containers(record.get("containers"))
    if not containers and "container" in record:
        containers = normalize_containers(record.get("container"))

    for container in containers:
        value = resolve_container_field(container, field)
        if value is None:
            continue
        if isinstance(value, str) and not value.strip():
            continue
        return value
    return None


def apply_match_highlighting_and_marker(
    sheet: Any,
    gt_col_index: int,
    predicted_col_index: int,
    match_marker_col_index: int,
) -> None:
    from openpyxl.styles import PatternFill  # type: ignore[import-not-found]

    match_fill = PatternFill(
        fill_type="solid",
        start_color="C6EFCE",
        end_color="C6EFCE",
    )

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        gt_value = normalize_value_for_match(row[gt_col_index].value)
        predicted_value = normalize_value_for_match(row[predicted_col_index].value)

        if gt_value is None or predicted_value is None:
            continue

        if gt_value == predicted_value:
            row[predicted_col_index].fill = match_fill
            row[match_marker_col_index].value = "literal_match"


def merge_predictions_into_validation_template(
    input_excel: Path,
    predictions_jsonl: Path,
    output_excel: Path,
) -> None:
    try:
        from openpyxl import load_workbook  # type: ignore[import-not-found]
    except ImportError as exc:
        logger.error(
            "Excel merge requires openpyxl. Install it with: pip install openpyxl"
        )
        raise RuntimeError(
            "Excel merge requires openpyxl. Install it with: pip install openpyxl"
        ) from exc

    if not input_excel.exists():
        logger.error("Template Excel file not found: %s", input_excel)
        raise FileNotFoundError(f"Template Excel file not found: {input_excel}")
    if not predictions_jsonl.exists():
        logger.error("Predictions JSONL file not found: %s", predictions_jsonl)
        raise FileNotFoundError(
            f"Predictions JSONL file not found: {predictions_jsonl}"
        )

    logger.info("Loading predictions from %s", predictions_jsonl)
    prediction_map, _ = load_prediction_records(predictions_jsonl)

    logger.info("Loading template Excel from %s", input_excel)
    workbook = load_workbook(filename=input_excel)
    sheet = workbook["Results"]

    if sheet.max_row < 1:
        logger.error("Excel file has no header row: %s", input_excel)
        raise ValueError(f"Excel file has no header row: {input_excel}")

    existing_headers = iter_header_values(
        next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    )

    required_headers = [
        "Id",
        "Container Duplicate",
        "Report",
        "GT Report Diagnosis",
        "Predicted Report Diagnosis",
        "GT Container",
        "Predicted Container",
        "Predicted Container Match",
        "GT Has Differential Diagnosis",
        "Predicted Has Differential Diagnosis",
        "Predicted Has Differential Diagnosis Match",
        "GT Is Lymph Node",
        "Predicted Is Lymph Node",
        "Predicted Is Lymph Node Match",
        "GT Is Definitive",
        "Predicted Is Definitive",
        "Predicted Is Definitive Match",
        "GT Has Prior Malignancy",
        "Predicted Has Prior Malignancy",
        "Predicted Has Prior Malignancy Match",
        "GT Has Concurrent Malignancy",
        "Predicted Has Concurrent Malignancy",
        "Predicted Has Concurrent Malignancy Match",
        "GT Container Diagnosis",
        "Predicted Container Diagnosis",
        "Predicted Container Dictionary Diagnosis",
        "Predicted Specimen",
        "Predicted Anatomic Location",
        "Predicted Container Diagnosis Code",
        "Predicted Code",
        "Predicted Dictionary Diagnosis",
        "Prediction Score",
        "Predicted Container Score",
    ]

    for header in required_headers:
        if header not in existing_headers:
            logger.error(
                "Required header '%s' not found in template. Found: %s",
                header,
                existing_headers,
            )
            raise ValueError(
                f"Required header '{header}' not found in template. Found: {existing_headers}"
            )

    header_indexes = {h: i for i, h in enumerate(existing_headers)}

    # Read all existing rows - keep them as-is
    base_rows: list[list[Any]] = []
    for row in sheet.iter_rows(
        min_row=2,
        max_row=sheet.max_row,
        min_col=1,
        max_col=len(existing_headers),
        values_only=True,
    ):
        base_rows.append(list(row))

    # Group rows by case_id
    rows_by_case_id: dict[int | str | None, list[tuple[int, list[Any]]]] = {}
    for row_num_offset, base_row in enumerate(base_rows, start=2):
        row_id_value = base_row[header_indexes["Id"]]
        normalized_id = normalize_case_id(row_id_value)
        if normalized_id not in rows_by_case_id:
            rows_by_case_id[normalized_id] = []
        rows_by_case_id[normalized_id].append((row_num_offset, base_row))

    # For each case_id, fill rows with containers from JSONL in order
    for normalized_id, rows_for_case in sorted(
        rows_by_case_id.items(),
        key=lambda x: x[1][-1][0],  # process bottom-up
        reverse=True,
    ):
        record = prediction_map.get(normalized_id)

        if record is None:
            continue

        containers = normalize_containers(record.get("containers"))

        num_excel_rows = len(rows_for_case)
        num_predicted = len(containers)

        # --- CASE 1: fill existing rows ---
        for row_offset, (row_num, base_row) in enumerate(rows_for_case):
            container = containers[row_offset] if row_offset < num_predicted else None

            # Case-level field (unchanged)
            sheet.cell(
                row=row_num,
                column=header_indexes["Predicted Report Diagnosis"] + 1,
                value=to_excel_value(record.get("primary_diagnosis")),
            )

            # New case-level fields
            sheet.cell(
                row=row_num,
                column=header_indexes["Predicted Code"] + 1,
                value=extract_primary_diagnosis_field(
                    record, "valid_primary_diagnosis_code"
                ),
            )
            sheet.cell(
                row=row_num,
                column=header_indexes["Prediction Score"] + 1,
                value=extract_primary_diagnosis_field(
                    record, "valid_primary_diagnosis_score"
                ),
            )

            sheet.cell(
                row=row_num,
                column=header_indexes["Predicted Dictionary Diagnosis"] + 1,
                value=extract_primary_diagnosis_field(
                    record, "valid_primary_diagnosis_name"
                ),
            )

            # New primary diagnosis group fields
            sheet.cell(
                row=row_num,
                column=header_indexes["Predicted Diagnosis Group 1"] + 1,
                value=extract_primary_diagnosis_field(
                    record, "valid_primary_diagnosis_group_1"
                ),
            )
            sheet.cell(
                row=row_num,
                column=header_indexes["Predicted Diagnosis Group 2"] + 1,
                value=extract_primary_diagnosis_field(
                    record, "valid_primary_diagnosis_group_2"
                ),
            )
            sheet.cell(
                row=row_num,
                column=header_indexes["Predicted Diagnosis Group 3"] + 1,
                value=extract_primary_diagnosis_field(
                    record, "valid_primary_diagnosis_group_3"
                ),
            )

            # --- Handle missing containers ---
            if container is None:
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container"] + 1,
                    value=None,
                )
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Is Lymph Node"] + 1,
                    value=None,
                )
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container Diagnosis"] + 1,
                    value=None,
                )
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Specimen"] + 1,
                    value=None,
                )
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Anatomic Location"] + 1,
                    value=None,
                )
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container Diagnosis Code"] + 1,
                    value=None,
                )
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container Score"] + 1,
                    value=None,
                )
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container Dictionary Diagnosis"]
                    + 1,
                    value=None,
                )
            else:
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container"] + 1,
                    value=resolve_container_label(container),
                )

                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Is Lymph Node"] + 1,
                    value=resolve_container_field(container, "is_lymph_node"),
                )

                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container Diagnosis"] + 1,
                    value=resolve_container_field(container, "diagnosis"),
                )
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Specimen"] + 1,
                    value=resolve_container_field(container, "specimen"),
                )

                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Anatomic Location"] + 1,
                    value=resolve_container_field(container, "anatomic_location"),
                )

                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container Diagnosis Code"] + 1,
                    value=extract_container_diagnosis_field(
                        container, "valid_diagnosis_code"
                    ),
                )
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container Score"] + 1,
                    value=extract_container_diagnosis_field(
                        container, "valid_diagnosis_score"
                    ),
                )
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container Dictionary Diagnosis"]
                    + 1,
                    value=extract_container_diagnosis_field(
                        container, "valid_diagnosis_name"
                    ),
                )

                # New container diagnosis group fields
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container Diagnosis Group 1"] + 1,
                    value=extract_container_diagnosis_field(
                        container, "valid_diagnosis_group_1"
                    ),
                )
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container Diagnosis Group 2"] + 1,
                    value=extract_container_diagnosis_field(
                        container, "valid_diagnosis_group_2"
                    ),
                )
                sheet.cell(
                    row=row_num,
                    column=header_indexes["Predicted Container Diagnosis Group 3"] + 1,
                    value=extract_container_diagnosis_field(
                        container, "valid_diagnosis_group_3"
                    ),
                )

            # Remaining unchanged fields
            sheet.cell(
                row=row_num,
                column=header_indexes["Predicted Has Differential Diagnosis"] + 1,
                value=to_excel_value(record.get("has_differential_diagnosis")),
            )

            sheet.cell(
                row=row_num,
                column=header_indexes["Predicted Is Definitive"] + 1,
                value=to_excel_value(record.get("is_definitive")),
            )

            sheet.cell(
                row=row_num,
                column=header_indexes["Predicted Has Prior Malignancy"] + 1,
                value=to_excel_value(record.get("has_prior_malignancy")),
            )

            sheet.cell(
                row=row_num,
                column=header_indexes["Predicted Has Concurrent Malignancy"] + 1,
                value=to_excel_value(record.get("has_concurrent_malignancy")),
            )

        # --- CASE 2: insert extra rows if predictions > excel rows ---
        if num_predicted > num_excel_rows:
            last_row_num = rows_for_case[-1][0]
            insert_at = last_row_num + 1

            extra_containers = containers[num_excel_rows:]

            for container in extra_containers:
                sheet.insert_rows(insert_at)

                new_row = insert_at

                # Id (copied from first row of this case block)
                sheet.cell(
                    row=new_row,
                    column=header_indexes["Id"] + 1,
                    value=rows_for_case[0][1][header_indexes["Id"]],
                )

                # Predicted Container
                sheet.cell(
                    row=new_row,
                    column=header_indexes["Predicted Container"] + 1,
                    value=resolve_container_label(container),
                )

                # Predicted Is Lymph Node
                sheet.cell(
                    row=new_row,
                    column=header_indexes["Predicted Is Lymph Node"] + 1,
                    value=resolve_container_field(container, "is_lymph_node"),
                )

                # Predicted Container Diagnosis
                sheet.cell(
                    row=new_row,
                    column=header_indexes["Predicted Container Diagnosis"] + 1,
                    value=resolve_container_field(container, "diagnosis"),
                )

                # Predicted Specimen
                sheet.cell(
                    row=new_row,
                    column=header_indexes["Predicted Specimen"] + 1,
                    value=resolve_container_field(container, "specimen"),
                )

                # Predicted Anatomic Location
                sheet.cell(
                    row=new_row,
                    column=header_indexes["Predicted Anatomic Location"] + 1,
                    value=resolve_container_field(container, "anatomic_location"),
                )

                # Predicted Container Diagnosis Code
                sheet.cell(
                    row=new_row,
                    column=header_indexes["Predicted Container Diagnosis Code"] + 1,
                    value=extract_container_diagnosis_field(
                        container, "valid_diagnosis_code"
                    ),
                )

                # Predicted Container Score
                sheet.cell(
                    row=new_row,
                    column=header_indexes["Predicted Container Score"] + 1,
                    value=extract_container_diagnosis_field(
                        container, "valid_diagnosis_score"
                    ),
                )

                # Predicted Container Dictionary Diagnosis
                sheet.cell(
                    row=new_row,
                    column=header_indexes["Predicted Container Dictionary Diagnosis"]
                    + 1,
                    value=extract_container_diagnosis_field(
                        container, "valid_diagnosis_name"
                    ),
                )

                # Predicted Container Diagnosis Group Fields
                sheet.cell(
                    row=new_row,
                    column=header_indexes["Predicted Container Diagnosis Group 1"] + 1,
                    value=extract_container_diagnosis_field(
                        container, "valid_diagnosis_group_1"
                    ),
                )
                sheet.cell(
                    row=new_row,
                    column=header_indexes["Predicted Container Diagnosis Group 2"] + 1,
                    value=extract_container_diagnosis_field(
                        container, "valid_diagnosis_group_2"
                    ),
                )
                sheet.cell(
                    row=new_row,
                    column=header_indexes["Predicted Container Diagnosis Group 3"] + 1,
                    value=extract_container_diagnosis_field(
                        container, "valid_diagnosis_group_3"
                    ),
                )

    # --- FIX Container Duplicate formulas after row insertions ---
    container_dup_col = header_indexes.get("Container Duplicate")
    id_col = header_indexes.get("Id")

    if container_dup_col is not None and id_col is not None:
        for row_idx in range(2, sheet.max_row + 1):
            id_value = sheet.cell(
                row=row_idx,
                column=id_col + 1,
            ).value

            # Skip rows that do not contain real data
            if id_value in (None, ""):
                continue

            sheet.cell(
                row=row_idx,
                column=container_dup_col + 1,
                value=f"=COUNTIF($A$2:A{row_idx},A{row_idx})=1",
            )

    match_pairs = [
        (
            "GT Report Diagnosis",
            "Predicted Dictionary Diagnosis",
            "Predicted Dictionary Diagnosis Match",
        ),
        ("GT Code", "Predicted Code", "Predicted Code Match"),
        (
            "GT Report Diagnosis Group 1",
            "Predicted Diagnosis Group 1",
            "Predicted Diagnosis Group 1 Match",
        ),
        (
            "GT Report Diagnosis Group 2",
            "Predicted Diagnosis Group 2",
            "Predicted Diagnosis Group 2 Match",
        ),
        (
            "GT Report Diagnosis Group 3",
            "Predicted Diagnosis Group 3",
            "Predicted Diagnosis Group 3 Match",
        ),
        ("GT Container", "Predicted Container", "Predicted Container Match"),
        (
            "GT Has Differential Diagnosis",
            "Predicted Has Differential Diagnosis",
            "Predicted Has Differential Diagnosis Match",
        ),
        (
            "GT Is Lymph Node",
            "Predicted Is Lymph Node",
            "Predicted Is Lymph Node Match",
        ),
        (
            "GT Is Definitive",
            "Predicted Is Definitive",
            "Predicted Is Definitive Match",
        ),
        (
            "GT Has Prior Malignancy",
            "Predicted Has Prior Malignancy",
            "Predicted Has Prior Malignancy Match",
        ),
        (
            "GT Has Concurrent Malignancy",
            "Predicted Has Concurrent Malignancy",
            "Predicted Has Concurrent Malignancy Match",
        ),
        (
            "GT Container Diagnosis Group 1",
            "Predicted Container Diagnosis Group 1",
            "Predicted Container Diagnosis Group 1 Match",
        ),
        (
            "GT Container Diagnosis Group 2",
            "Predicted Container Diagnosis Group 2",
            "Predicted Container Diagnosis Group 2 Match",
        ),
        (
            "GT Container Diagnosis Group 3",
            "Predicted Container Diagnosis Group 3",
            "Predicted Container Diagnosis Group 3 Match",
        ),
    ]

    for gt_header, predicted_header, match_header in match_pairs:
        gt_col_index = header_indexes.get(gt_header)
        predicted_col_index = header_indexes.get(predicted_header)
        match_marker_col_index = header_indexes.get(match_header)

        if (
            gt_col_index is None
            or predicted_col_index is None
            or match_marker_col_index is None
        ):
            continue

        apply_match_highlighting_and_marker(
            sheet=sheet,
            gt_col_index=gt_col_index,
            predicted_col_index=predicted_col_index,
            match_marker_col_index=match_marker_col_index,
        )

    # Find last row with data
    id_col = header_indexes["Id"] + 1

    last_data_row = 1
    for row_idx in range(sheet.max_row, 1, -1):
        if sheet.cell(row=row_idx, column=id_col).value not in (None, ""):
            last_data_row = row_idx
            break

    # Expand Excel table to actual data rows
    for table in sheet.tables.values():
        start_cell, end_cell = table.ref.split(":")

        start_col = start_cell.rstrip("0123456789")
        start_row = int(start_cell[len(start_col) :])

        end_col = end_cell.rstrip("0123456789")

        table.ref = f"{start_col}{start_row}:{end_col}{last_data_row}"

    output_excel.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_excel)
    logger.info("Wrote merged Excel to %s", output_excel)


def load_config(config_path: str) -> dict:
    """Load configuration from YAML file."""
    config_file = Path(config_path)
    if not config_file.exists():
        logger.error("Configuration file not found: %s", config_path)
        raise FileNotFoundError(f"Configuration file not found: {config_path}")

    with open(config_file) as f:
        config = yaml.safe_load(f)

    logger.info("Loaded configuration from %s", config_path)

    return config


def argparse_setup():
    parser = argparse.ArgumentParser(
        description="Merge predictions into validation template",
    )
    parser.add_argument(
        "--config",
        default="configs/excel_merge.yaml",
        help="Path to Excel merge configuration YAML file",
    )

    args = parser.parse_args()

    return args.config


def main() -> int:
    """Main entry point for Excel merge with config loading."""

    config_path = argparse_setup()

    config = load_config(config_path)

    merge_predictions_into_validation_template(
        input_excel=Path(config.get("input_excel")),
        predictions_jsonl=Path(config.get("predictions_jsonl")),
        output_excel=Path(config.get("output_excel")),
    )
    logger.info("Wrote merged Excel: %s", config.get("output_excel"))
    return 0
